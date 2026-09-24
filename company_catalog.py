"""Import a local JPX monthly workbook; never fetch or call an LLM implicitly."""
import argparse
from contextlib import closing
from datetime import datetime,timezone
import hashlib
import json
from pathlib import Path
import re
import unicodedata
from system_data import encode,digest,stamp,readonly

SOURCE='https://www.jpx.co.jp/markets/statistics-equities/misc/01.html'
# Broad sector hints only, never proof of a company's specific business.
SECTORS={'銀行業':['banking'],'不動産業':['real_estate'],'医薬品':['healthcare'],
    '電気・ガス業':['energy'],'石油・石炭製品':['energy'],'鉱業':['energy'],
    '電気機器':['semiconductor'],'精密機器':['semiconductor'],'機械':['semiconductor'],
    '輸送用機器':['export_manufacturing'],'卸売業':['trade']}
MARKETS={'プライム（内国株式）','スタンダード（内国株式）','グロース（内国株式）'}


def parse_rows(rows,now):
    required={'日付','コード','銘柄名','市場・商品区分','33業種区分'}
    output=[];seen=set();dates=set()
    for row in rows:
        if not required<=set(row):raise ValueError('JPX required columns missing')
        if row['市場・商品区分'] not in MARKETS:continue
        code=str(row['コード']).removesuffix('.0').upper()
        if re.fullmatch(r'[0-9A-Z]{5}',code):continue  # Preferred/class shares are outside 4-character universe.
        raw_day=str(row['日付']).removesuffix('.0')
        if not re.fullmatch(r'[0-9A-Z]{4}',code) or code in seen:raise ValueError('Invalid/duplicate code')
        day=datetime.strptime(raw_day,'%Y%m%d').date().isoformat()
        if day>now.date().isoformat():raise ValueError('Future listing date')
        name=str(row['銘柄名']).strip();industry=str(row['33業種区分']).strip()
        if not name or not industry or industry=='-':raise ValueError('Missing company/sector')
        output.append(dict(symbol=code,name=name,industry=industry,market=row['市場・商品区分'],asof=day))
        seen.add(code);dates.add(day)
    if not output or len(dates)!=1:raise ValueError('Empty/mixed dated JPX snapshot')
    return sorted(output,key=lambda x:x['symbol'])


def read_file(path,now):
    path=Path(path)
    if path.suffix.lower()=='.csv':
        import csv
        with path.open(encoding='utf-8-sig',newline='') as f:rows=list(csv.DictReader(f))
    elif path.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=True)
        try:
            iterator=book.worksheets[0].iter_rows(values_only=True);headers=next(iterator)
            rows=[dict(zip(headers,row)) for row in iterator if any(x is not None for x in row)]
        finally:book.close()
    elif path.suffix.lower()=='.xls':
        import xlrd
        book=xlrd.open_workbook(path);sheet=book.sheet_by_index(0)
        headers=sheet.row_values(0)
        rows=[dict(zip(headers,sheet.row_values(i))) for i in range(1,sheet.nrows)]
    else:raise ValueError('Use original JPX .xls/.xlsx or UTF-8 CSV with identical columns')
    return parse_rows(rows,now)


def import_file(db_path,path,now=None):
    now=now or datetime.now(timezone.utc);records=read_file(path,now)
    from company_graph import connect
    raw_hash=hashlib.sha256(Path(path).read_bytes()).hexdigest()
    payload=dict(source=SOURCE,source_date=records[0]['asof'],file_sha256=raw_hash,records=records)
    ident=digest(payload)
    with closing(connect(db_path)) as db,db:
        db.execute('CREATE TABLE IF NOT EXISTS listing_snapshots(id TEXT PRIMARY KEY,recorded_at TEXT,payload TEXT)')
        previous=db.execute('SELECT payload FROM listing_snapshots ORDER BY recorded_at DESC,rowid DESC LIMIT 1').fetchone()
        if previous and json.loads(previous[0])['source_date']>payload['source_date']:
            raise ValueError('Older listing snapshot cannot replace the latest master')
        db.execute('INSERT OR IGNORE INTO listing_snapshots VALUES (?,?,?)',(ident,now.isoformat(),encode(payload)))
    return dict(companies=len(records),source_date=payload['source_date'],file_sha256=raw_hash,source=SOURCE)


def latest(path,asof):
    if not Path(path).is_file():return None
    with closing(readonly(path)) as db:
        if not db.execute("SELECT 1 FROM sqlite_master WHERE name='listing_snapshots'").fetchone():return None
        for ident,recorded,raw in db.execute('SELECT * FROM listing_snapshots ORDER BY recorded_at DESC,rowid DESC'):
            if stamp(recorded)>asof:continue
            payload=json.loads(raw)
            if digest(payload)!=ident:raise ValueError('Listing snapshot integrity mismatch')
            return dict(payload,revision=ident,recorded_at=recorded)


def merge(registry,snapshot):
    if not snapshot:return registry
    existing={r['symbol']:r for r in registry};out=[]
    for listed in snapshot['records']:
        old=existing.get(listed['symbol'])
        c=dict(old) if old else dict(symbol=listed['symbol'],name=listed['name'],aliases=[],
            industry=listed['industry'],business='主要事業未確認。業種分類から個別事業を断定しません。',
            updated_at=snapshot['recorded_at'],relations=[],recorded_at=snapshot['recorded_at'],revision='')
        c['industry']=listed['industry']
        # Name changes require re-verifying old business evidence, not silently inheriting it.
        if old and unicodedata.normalize('NFKC',old['name'])!=unicodedata.normalize('NFKC',listed['name']):
            c.update(name=listed['name'],aliases=[],business='名称変更・表記差異のため事業情報の再確認が必要',relations=[])
        c['relations']=list(c['relations'])
        known={r['topic'] for r in c['relations']}
        for topic in SECTORS.get(listed['industry'],[]):
            if topic not in known:
                c['relations'].append(dict(topic=topic,role='不明',business='業種からの広い探索候補。個別事業は未確認',
                    state='推定',direct_business='不明',source=SOURCE,quote=listed['industry'],verified_at=None,
                    revenue_ratio=None,profit_ratio=None,ratio_evidence=''))
        c['listing']=dict(listed,source=SOURCE,file_sha256=snapshot['file_sha256'],recorded_at=snapshot['recorded_at'])
        c['revision']=digest([c['revision'],snapshot['revision'],listed])
        c['recorded_at']=max(stamp(c['recorded_at']),stamp(snapshot['recorded_at'])).isoformat()
        out.append(c)
    return out


def main(argv=None):
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('command',choices=['import-jpx','status','template'])
    p.add_argument('--file',type=Path);p.add_argument('--symbol');p.add_argument('--output',type=Path)
    p.add_argument('--config',type=Path,default=Path(__file__).resolve().parent/'config/paper.json')
    a=p.parse_args(argv)
    from system_runtime import read_config
    from news_discovery import settings
    from company_graph import companies
    c=read_config(a.config);path=settings(c)['companies_db'];now=datetime.now(timezone.utc)
    if a.command=='import-jpx':
        if not a.file:p.error('--file required')
        print(encode(import_file(path,a.file,now)));return
    rows=companies(path,now)
    if a.command=='status':
        print(encode(dict(companies=len(rows),confirmed_business=sum(any(r['state']=='確認済み' for r in x['relations']) for x in rows),
            sector_hint_only=sum(bool(x['relations']) and all(r['state']!='確認済み' for r in x['relations']) for x in rows),
            listing_source_date=(latest(path,now) or {}).get('source_date'))));return
    if not a.symbol or not a.output:p.error('--symbol and new --output required')
    row=next((x for x in rows if x['symbol']==a.symbol),None)
    if row is None:raise ValueError('Unknown symbol')
    clean={k:row[k] for k in ('symbol','name','aliases','industry','business','updated_at','relations')}
    clean['updated_at']=now.isoformat()
    with a.output.open('x',encoding='utf-8') as f:f.write(encode([clean]))


if __name__=='__main__':main()
